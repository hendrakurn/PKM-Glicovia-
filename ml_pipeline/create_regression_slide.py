from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUT = "Regression_Model_Comparison.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

WHITE = RGBColor(255,255,255); NAVY = RGBColor(22,43,72); BLUE = RGBColor(225,239,255)
GREEN = RGBColor(224,245,231); GRID = RGBColor(207,216,226); TEXT = RGBColor(31,41,55)
MUTED = RGBColor(86,101,118); ACCENT = RGBColor(35,94,160)

def box(x,y,w,h,fill=None,line=None,radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill or WHITE
    shape.line.color.rgb = line or fill or WHITE
    if radius:
        shape.adjustments[0] = 0.08
    return shape

def text(x,y,w,h,s,size=12,color=TEXT,bold=False,align=PP_ALIGN.LEFT,font="Aptos",valign=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.margin_left = Inches(.08); tf.margin_right = Inches(.08)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

box(0,0,13.333,7.5,WHITE)
box(.55,.48,.12,.72,ACCENT)
text(.82,.38,11.9,.42,"Regression Model Comparison",25,NAVY,True)
text(.83,.86,11.6,.26,"Cross-validation performance and hold-out evaluation",10,MUTED)

cols = [("Model",3.25),("CV R² Mean",1.25),("Test R²",1.05),("Test RMSE",1.15),("Test MAE",1.05),("Result",3.05)]
x0=.7; y0=1.38; header_h=.43; row_h=.47
box(x0,y0,11.93,header_h,NAVY,NAVY)
x=x0
for label,w in cols:
    text(x,y0+.02,w,header_h-.04,label,10,WHITE,True,PP_ALIGN.CENTER if label!="Model" and label!="Result" else PP_ALIGN.LEFT)
    x+=w

rows = [
("Ridge","0.9263","0.9777","8.1407","6.6821","Best test R²",BLUE),
("LinearRegression","0.9261","0.9777","8.1433","6.7469","Strong baseline",WHITE),
("TunedSVR_RBF","0.9568","0.9689","9.6212","7.3927","Recommended by CV",GREEN),
("GradientBoostingRegressor","0.8958","0.9687","9.6557","7.0889","Strong nonlinear model",WHITE),
("RandomForestRegressor","0.8980","0.9442","12.8812","9.3499","Lower performance",WHITE),
("KNeighborsRegressor","0.9167","0.9440","12.9039","8.9102","Lower performance",WHITE),
("SVR_RBF","0.7848","0.8825","18.6958","11.6415","Untuned SVM",WHITE),
("DummyRegressor","-0.0684","-0.0001","54.5510","41.3773","Baseline only",WHITE),]
for i,row in enumerate(rows):
    y=y0+header_h+i*row_h
    model,cv,tr,rmse,mae,result,fill=row
    box(x0,y,11.93,row_h,fill,GRID)
    vals=[model,cv,tr,rmse,mae,result]; x=x0
    for j,((label,w),val) in enumerate(zip(cols, vals)):
        text(x,y+.015,w,row_h-.03,val,9.2,TEXT,j==0,PP_ALIGN.CENTER if j in [1,2,3,4] else PP_ALIGN.LEFT)
        x+=w

text(.72,5.48,1.15,.22,"CONCLUSION",9,ACCENT,True)
box(.7,5.75,11.93,.91,RGBColor(247,249,252),RGBColor(223,230,238),True)
conclusion=("TunedSVR_RBF is the recommended model because it achieved the highest cross-validation R², providing the strongest average "
            "generalization across folds. Ridge achieved the highest test R² on the single hold-out split.")
text(.92,5.88,11.5,.58,conclusion,11,NAVY,False)
text(.72,7.08,11.9,.18,"Note: R² indicates explained variance; RMSE and MAE are reported in the target variable’s units.",8.5,MUTED)

prs.save(OUT)
print(OUT)
